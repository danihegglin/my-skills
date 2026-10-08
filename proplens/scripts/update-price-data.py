"""Rebuilds src/data/prices.json from official Swiss open data.

    pip install openpyxl
    python3 scripts/update-price-data.py

Sources (all open data, linked from opendata.swiss):
- Swiss Federal Statistical Office (BFS), structural survey: average net rent per month
  (T 09.03.03.01) and per m2 (T 09.03.03.05) by canton and number of rooms, rent per m2 by
  tenancy duration (T 09.03.03.63), by construction period (T 09.03.03.06), by urban/rural
  municipality type (T 09.03.03.08) and for the ten largest cities (T 09.03.03.42, three years
  pooled), latest edition of each.
- Statistical Office of the Canton of Zurich: prices of condominiums and single-family houses
  sold, by property-market region, and the assignment of municipalities to those regions.
- City of Winterthur housing monitor: advertised rents (listings) against rents tenants pay,
  by number of rooms. The only open comparison of the two in Switzerland.
"""

import csv
import io
import json
import sys
import time
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl

BFS = "https://dam-api.bfs.admin.ch/hub/api/dam/assets/{}/master"
RENT_MONTHLY = 36398431  # T 09.03.03.01
RENT_M2 = 36398449  # T 09.03.03.05
RENT_TENURE = 36398526  # T 09.03.03.63
RENT_PERIOD = 36398457  # T 09.03.03.06
RENT_TYPE = 20784959  # T 09.03.03.08
RENT_CITIES = 36506661  # T 09.03.03.42, 2022-2024 pooled
# Official BFS municipality numbers of the ten largest cities, by their name in that table.
CITY_NUMBERS = {
    "Zürich": 261, "Genf": 6621, "Basel": 2701, "Lausanne": 5586, "Bern": 351,
    "Winterthur": 230, "Luzern": 1061, "St.Gallen": 3203, "Lugano": 5192, "Biel": 371,
}
CITY_CANTONS = {261: "ZH", 6621: "GE", 2701: "BS", 5586: "VD", 351: "BE", 230: "ZH", 1061: "LU", 3203: "SG", 5192: "TI", 371: "BE"}
ZH = "https://daten.statistik.zh.ch/ogd/daten/ressourcen/KTZH_00003158_0000{}"
ZH_CONDOS = ZH.format("6784.csv")
ZH_HOUSES = ZH.format("6781.csv")
ZH_REGIONS = ZH.format("6783.xlsx")
WINTERTHUR_RENTS = "https://daten.statistik.zh.ch/ogd/daten/ressourcen/KTZH_00003156_00006770.csv"

# Cantons in the BFS's official order, as the rows of every table above follow it.
CANTONS = "ZH BE LU UR SZ OW NW GL ZG FR SO BS BL SH AR AI SG GR AG TG TI VD VS NE GE JU".split()
ROOMS = ["total", "1", "2", "3", "4", "5", "6+"]
# Region names in the municipality list → names used in the price tables.
REGION_NAMES = {
    "Land": "Landgemeinden",
    "Agglomeration": "Agglomerationsgemeinden",
    "See": "Seegemeinden",
    "Stadt Zürich": "Stadt Zürich",
    "Stadt Winterthur": "Stadt Winterthur",
}

OUT = Path(__file__).resolve().parent.parent / "src" / "data" / "prices.json"


def fetch(url: str, attempts: int = 4) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "PropLens data update"})
    for attempt in range(attempts):
        try:
            with urllib.request.urlopen(req, timeout=60) as res:
                return res.read()
        except OSError:
            if attempt == attempts - 1:
                raise
            time.sleep(2 ** attempt)
    raise AssertionError("unreachable")


def latest_sheet(asset: int):
    wb = openpyxl.load_workbook(io.BytesIO(fetch(BFS.format(asset))), read_only=True, data_only=True)
    year = max(int(name) for name in wb.sheetnames if name.isdigit())
    return year, list(wb[str(year)].iter_rows(values_only=True))


def numeric_rows(rows):
    """Data rows: a label followed by (value, ± confidence) pairs for total and 1…6+ rooms."""
    out = []
    for row in rows:
        cells = [c for c in row if c is not None]
        if len(cells) >= 1 + 2 * len(ROOMS) and isinstance(cells[1], (int, float)):
            out.append((str(cells[0]).strip(), cells[1 : 1 + 2 * len(ROOMS)]))
    return out


def number(v):
    """Cells are numbers, "(123)" for extrapolations from few cases, or "X" when suppressed."""
    if isinstance(v, (int, float)):
        return v
    text = str(v).strip().strip("()")
    try:
        return float(text)
    except ValueError:
        return None


def pairs(values):
    return {room: [number(values[2 * i]), number(values[2 * i + 1])] for i, room in enumerate(ROOMS)}


def by_canton(rows):
    data = numeric_rows(rows)
    if len(data) < 27:
        sys.exit(f"Expected Switzerland + 26 cantons, found {len(data)} rows")
    result = {"CH": {"name": data[0][0], **{"values": pairs(data[0][1])}}}
    for code, (name, values) in zip(CANTONS, data[1:27]):
        result[code] = {"name": name, "values": pairs(values)}
    return result


def first_row_ratios(rows, labels):
    """Ratios of each column to the table's total, from the first (Switzerland) data row."""
    for row in rows:
        cells = [c for c in row if c is not None]
        if len(cells) >= 1 + 2 * len(labels) and isinstance(cells[1], (int, float)):
            values = cells[1 : 1 + 2 * len(labels) : 2]
            return {label: round(v / values[0], 3) for label, v in zip(labels, values)}
    sys.exit("No data row found")


def type_ratios(rows):
    """Urban / intermediate / rural rent per m2 relative to Switzerland, by rooms."""
    data = numeric_rows(rows)
    swiss = dict(zip(ROOMS, data[0][1][::2]))
    out = {}
    for key, (_, values) in zip(["urban", "intermediate", "rural"], data[1:4]):
        out[key] = {room: round(v / swiss[room], 3) for room, v in zip(ROOMS, values[::2])}
    return out


def city_factors(cantons):
    """Each city's monthly rent relative to its canton, both measured against Switzerland."""
    wb = openpyxl.load_workbook(io.BytesIO(fetch(BFS.format(RENT_CITIES))), read_only=True, data_only=True)
    sheet = wb.sheetnames[0]
    data = numeric_rows(wb[sheet].iter_rows(values_only=True))
    swiss = dict(zip(ROOMS, data[0][1][::2]))
    out = {}
    for name, values in data[1:]:
        number = CITY_NUMBERS.get(name)
        if number is None:
            sys.exit(f"Unknown city in rent table: {name}")
        canton = cantons[CITY_CANTONS[number]]["monthly"]
        swiss_now = cantons["CH"]["monthly"]
        out[str(number)] = {
            "name": name,
            "factors": {room: round((v / swiss[room]) / (canton[room][0] / swiss_now[room][0]), 3) for room, v in zip(ROOMS, values[::2])},
        }
    return sheet, out


def asking_premium():
    """Advertised ÷ paid rent, city-wide in Winterthur, by rooms (all owner types)."""
    rows = list(csv.DictReader(io.StringIO(fetch(WINTERTHUR_RENTS).decode("utf-8-sig"))))
    city = [r for r in rows if r["geom_code"] == "gemeinde" and r["eigentuemer_typ_nr"] == "0" and r["miete_monat_netto_avg"]]
    latest = max(r["jahr_bis"] for r in city)
    city = [r for r in city if r["jahr_bis"] == latest]
    by = {(r["typ"], r["whg_zim_group_nr"]): float(r["miete_monat_netto_avg"]) for r in city}
    factors = {}
    for i, room in enumerate(ROOMS[1:], start=1):
        asked, paid = by.get(("Angebotsmieten", str(i))), by.get(("Bestandsmieten", str(i)))
        if asked and paid:
            factors[room] = round(asked / paid, 3)
    years = f"{city[0]['jahr_von']}-{latest}"
    return years, factors


def zurich_sales(url: str):
    rows = list(csv.DictReader(io.StringIO(fetch(url).decode("utf-8-sig"))))
    year = max(int(r["Jahr"]) for r in rows)
    out = {}
    for r in rows:
        if int(r["Jahr"]) != year:
            continue
        out[r["region"]] = {k: round(float(r[v])) for k, v in [("n", "Verkaeufe"), ("q10", "Q10"), ("q25", "Q25"), ("median", "Median"), ("q75", "Q75"), ("q90", "Q90")]}
    return year, out


def zurich_regions():
    """The workbook's drawings are broken for openpyxl, so read the sheet XML directly."""
    z = zipfile.ZipFile(io.BytesIO(fetch(ZH_REGIONS)))
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    strings = ["".join(t.text or "" for t in si.iter(f"{{{ns['m']}}}t")) for si in ET.fromstring(z.read("xl/sharedStrings.xml")).findall("m:si", ns)]
    names = [s.get("name") for s in ET.fromstring(z.read("xl/workbook.xml")).find("m:sheets", ns)]
    sheet = ET.fromstring(z.read(f"xl/worksheets/sheet{names.index('Gemeindeliste') + 1}.xml"))
    mapping = {}
    for row in sheet.iter(f"{{{ns['m']}}}row"):
        cells = []
        for c in row.findall("m:c", ns):
            v = c.find("m:v", ns)
            if v is not None:
                cells.append(strings[int(v.text)] if c.get("t") == "s" else v.text)
        if len(cells) == 3 and cells[0].strip().isdigit():
            mapping[str(int(cells[0]))] = REGION_NAMES[cells[2].strip()]
    return mapping


def main():
    year, monthly_rows = latest_sheet(RENT_MONTHLY)
    year_m2, m2_rows = latest_sheet(RENT_M2)
    year_tenure, tenure_rows = latest_sheet(RENT_TENURE)
    if not year == year_m2 == year_tenure:
        sys.exit(f"Rent tables disagree on the latest year: {year}, {year_m2}, {year_tenure}")

    monthly = by_canton(monthly_rows)
    per_m2 = by_canton(m2_rows)
    cantons = {
        code: {"name": monthly[code]["name"], "monthly": monthly[code]["values"], "perM2": per_m2[code]["values"]}
        for code in monthly
    }

    # Rows: all tenancies, then moved in < 2 years ago into a new build, then into an older building, …
    tenure = numeric_rows(tenure_rows)
    total, new_build, existing = (dict(zip(ROOMS, values[::2])) for _, values in tenure[:3])
    new_lease = {
        "newBuild": {room: round(new_build[room] / total[room], 3) for room in ROOMS},
        "existing": {room: round(existing[room] / total[room], 3) for room in ROOMS},
    }

    year_period, period_rows = latest_sheet(RENT_PERIOD)
    periods = first_row_ratios(
        period_rows,
        ["total", "<1919", "1919-1945", "1946-1960", "1961-1970", "1971-1980", "1981-1990", "1991-2000", "2001-2010", "2011-2020", "2021+"],
    )
    year_type, type_rows = latest_sheet(RENT_TYPE)
    city_years, cities = city_factors(cantons)

    asking_years, asking = asking_premium()

    condo_year, condos = zurich_sales(ZH_CONDOS)
    house_year, houses = zurich_sales(ZH_HOUSES)

    data = {
        "rent": {
            "year": year,
            "source": "Swiss Federal Statistical Office (BFS), structural survey",
            "note": "Average net rent (excluding service charges) of occupied rented flats; ± is the 95% confidence interval.",
            "cantons": cantons,
            "newLease": new_lease,
            "periods": {"year": year_period, "factors": periods},
            "municipalityTypes": {"year": year_type, "factors": type_ratios(type_rows)},
            "cities": {"years": city_years, "byMunicipality": cities},
            "askingPremium": {"years": asking_years, "source": "City of Winterthur housing monitor", "factors": asking},
        },
        "zurichSales": {
            "year": min(condo_year, house_year),
            "source": "Statistical Office of the Canton of Zurich, property transfer statistics",
            "condos": condos,
            "houses": houses,
            "municipalities": zurich_regions(),
        },
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"Wrote {OUT} (rents {year}, Zurich sales {data['zurichSales']['year']}, {len(data['zurichSales']['municipalities'])} municipalities)")


if __name__ == "__main__":
    main()
